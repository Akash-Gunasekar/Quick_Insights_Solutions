import fitz  # PyMuPDF
from neo4j import GraphDatabase
import win32com.client as win32
from docx import Document
import os
import openpyxl

# Function to extract text from .txt files
def extract_text_from_txt(txt_path):
    with open(txt_path, 'r', encoding='utf-8') as file:
        text = file.read()
    return text

# Extract text from PDF
def extract_text_from_pdf(pdf_path):
    document = fitz.open(pdf_path)
    text = ""
    for page_num in range(len(document)):
        page = document.load_page(page_num)
        text += page.get_text()
    return text

# Function to extract text from Excel files
def extract_text_from_excel(excel_path):
    workbook = openpyxl.load_workbook(excel_path)
    text = ""
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value:
                    text += str(cell.value) + " "
            text += "\n"
    return text

# Function to extract text from .docx files
def extract_text_from_docx(docx_path):
    doc = Document(docx_path)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

# Function to extract text from .doc files
def extract_text_from_doc(doc_path):
    word = win32.Dispatch("Word.Application")
    doc = word.Documents.Open(doc_path)
    text = doc.Content.Text
    doc.Close()
    word.Quit()
    return text

def create_relationship(session, title, i, relationship_title):
    query = f"""
    MATCH (d:Document {{title: $title}}), (p:Paragraph {{id: $id, document_title: $title}})
    CREATE (d)-[:{relationship_title}]->(p)
    """
    session.run(query, title=title, id=i)

# Specific function to load text from a PDF into Neo4j
def load_pdf_to_neo4j(uri, user, password, file_path):
    text = extract_text_from_pdf(file_path)
    title = os.path.basename(file_path)
    relationship_title = "Contains"
    driver = GraphDatabase.driver(uri, auth=(user, password))
    with driver.session() as session:
        # Create a root node for the document
        session.run("CREATE (d:Document {title: $title})", title=title)

        # Split the text into paragraphs or lines
        paragraphs = text.split('\n\n')
        for i, para in enumerate(paragraphs):
            # Create a node for each paragraph
            session.run(
                "CREATE (p:Paragraph {id: $id, content: $content, document_title: $title})",
                id=i, content=para.strip(), title=title
            )
            # Create a relationship to the document with a title
            create_relationship(session, title, i, relationship_title)
    driver.close()

# Specific function to load text from a .txt file into Neo4j
def load_txt_to_neo4j(uri, user, password, file_path):
    text = extract_text_from_txt(file_path)
    title = os.path.basename(file_path)
    relationship_title = "Contains"
    driver = GraphDatabase.driver(uri, auth=(user, password))
    with driver.session() as session:
        # Create a root node for the document
        session.run("CREATE (d:Document {title: $title})", title=title)

        # Split the text into paragraphs or lines
        paragraphs = text.split('\n\n')
        for i, para in enumerate(paragraphs):
            # Create a node for each paragraph
            session.run(
                "CREATE (p:Paragraph {id: $id, content: $content, document_title: $title})",
                id=i, content=para.strip(), title=title
            )
            # Create a relationship to the document with a title
            create_relationship(session, title, i, relationship_title)
    driver.close()

# Specific function to load text from an Excel file into Neo4j
def load_excel_to_neo4j(uri, user, password, file_path):
    text = extract_text_from_excel(file_path)
    title = os.path.basename(file_path)
    relationship_title = "Contains"
    driver = GraphDatabase.driver(uri, auth=(user, password))
    with driver.session() as session:
        # Create a root node for the document
        session.run("CREATE (d:Document {title: $title})", title=title)

        # Split the text into paragraphs or lines
        paragraphs = text.split('\n\n')
        for i, para in enumerate(paragraphs):
            # Create a node for each paragraph
            session.run(
                "CREATE (p:Paragraph {id: $id, content: $content, document_title: $title})",
                id=i, content=para.strip(), title=title
            )
            # Create a relationship to the document with a title
            create_relationship(session, title, i, relationship_title)
    driver.close()

# Specific function to load text from a .docx file into Neo4j
def load_docx_to_neo4j(uri, user, password, file_path):
    text = extract_text_from_docx(file_path)
    title = os.path.basename(file_path)
    relationship_title = "Contains"
    driver = GraphDatabase.driver(uri, auth=(user, password))
    with driver.session() as session:
        # Create a root node for the document
        session.run("CREATE (d:Document {title: $title})", title=title)

        # Split the text into paragraphs or lines
        paragraphs = text.split('\n\n')
        for i, para in enumerate(paragraphs):
            # Create a node for each paragraph
            session.run(
                "CREATE (p:Paragraph {id: $id, content: $content, document_title: $title})",
                id=i, content=para.strip(), title=title
            )
            # Create a relationship to the document with a title
            create_relationship(session, title, i, relationship_title)
    driver.close()

# Specific function to load text from a .doc file into Neo4j
def load_doc_to_neo4j(uri, user, password, file_path):
    text = extract_text_from_doc(file_path)
    title = os.path.basename(file_path)
    relationship_title = "Contains"
    driver = GraphDatabase.driver(uri, auth=(user, password))
    with driver.session() as session:
        # Create a root node for the document
        session.run("CREATE (d:Document {title: $title})", title=title)

        # Split the text into paragraphs or lines
        paragraphs = text.split('\n\n')
        for i, para in enumerate(paragraphs):
            # Create a node for each paragraph
            session.run(
                "CREATE (p:Paragraph {id: $id, content: $content, document_title: $title})",
                id=i, content=para.strip(), title=title
            )
            # Create a relationship to the document with a title
            create_relationship(session, title, i, relationship_title)
    driver.close()

# Dictionary to map file extensions to loading functions
loaders = {
    '.pdf': load_pdf_to_neo4j,
    '.xlsx': load_excel_to_neo4j,
    '.docx': load_docx_to_neo4j,
    '.doc': load_doc_to_neo4j,
    '.txt': load_txt_to_neo4j
}

if __name__ == "__main__":
    # List of file paths
    file_paths = [
        "uploads/detailed_document.pdf",
        "uploads/transaction_doc.txt",
        "uploads/transaction01_doc.txt",

    ]
    #file_paths = ["detailed_document.pdf","example_document.docx","sample_spreadsheet.xlsx","text_file.txt","another_document.pdf","yet_another_document.pdf"]
    
    # Neo4j connection details
    uri = "bolt://localhost:7687"
    user = "neo4j"
    password = "Password@24"

    for file_path in file_paths:
        # Get the file extension
        file_extension = os.path.splitext(file_path)[1].lower()

        # Check if the extension is supported
        if file_extension in loaders:
            # Load text and insert into Neo4j using the appropriate function
            load_function = loaders[file_extension]
            load_function(uri, user, password, file_path)

            print(f"Data from {file_path} has been inserted into Neo4j successfully.")
        else:
            print(f"Unsupported file extension: {file_extension}")

    # Create relationships between different documents
    driver = GraphDatabase.driver(uri, auth=(user, password))
    with driver.session() as session:
        # Example of creating relationships between documents
        session.run("""
            MATCH (d1:Document {title: 'detailed_document.pdf'}),
                  (d2:Document {title: 'transaction_doc.txt'}),
                  (d3:Document {title: 'transaction01_doc.txt'})
                  
             CREATE (d1)-[:RELATED_TO {relationship_title: 'Contains'}]->(d2),
                   (d2)-[:RELATED_TO {relationship_title: 'Contains'}]->(d3),
                   (d3)-[:RELATED_TO {relationship_title: 'Contains'}]->(d1)
                
        """)
    driver.close()

    print("Relationships between documents have been created successfully.")